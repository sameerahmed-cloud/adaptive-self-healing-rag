import os
import re
import json
import csv
import hashlib
import shutil
import qdrant_client
import logging
import warnings
import uuid
import time
import math
from collections import Counter, defaultdict

from pathlib import Path
from dataclasses import dataclass
from typing import List, Dict, Optional, Tuple
from json_repair import repair_json

from dotenv import load_dotenv

from llama_index.core import (
    Settings,
    Document,
    VectorStoreIndex,
    SummaryIndex,
    StorageContext,
    load_index_from_storage,
)

from llama_index.core.callbacks import (
    CallbackManager, TokenCountingHandler,
)

from llama_index.vector_stores.qdrant import QdrantVectorStore

from llama_index.core.node_parser import (
    SentenceSplitter,
    MarkdownNodeParser,
    CodeSplitter,
)

from llama_index.readers.file import HTMLTagReader

from llama_index.core.ingestion import IngestionPipeline

from llama_index.core.query_engine import RouterQueryEngine

from llama_index.core.selectors import LLMSingleSelector

from llama_index.core.tools import (
    QueryEngineTool,
    ToolMetadata,
)

from llama_index.core.schema import NodeRelationship, RelatedNodeInfo, NodeWithScore
from llama_index.core.postprocessor import PrevNextNodePostprocessor

from llama_index.llms.google_genai import GoogleGenAI

from llama_index.embeddings.huggingface import (
    HuggingFaceEmbedding,
)

from llama_parse import LlamaParse

logging.getLogger("qdrant_client").setLevel(logging.WARNING)
logging.getLogger("httpx").setLevel(logging.WARNING)
logging.getLogger("llama_index").setLevel(logging.WARNING)


warnings.filterwarnings("ignore", category=DeprecationWarning)

class APIUsageTracker:
    """Lightweight, real-time observability for the current application session."""

    def __init__(self):
        self.started_at = time.perf_counter()
        self.llm_token_counter = TokenCountingHandler()
        self.callback_manager = CallbackManager([self.llm_token_counter])

        # External API / model activity
        self.llm_calls = 0
        self.llm_operations = {}
        self.llama_parse_calls = 0
        self.llama_parse_duration = 0.0

        # Local pipeline activity
        self.embedding_calls = 0
        self.embedding_tokens = 0
        self.embedding_duration = 0.0
        self.nodes_indexed = 0
        self.indexing_duration = 0.0
        self.chunking_duration = 0.0

        # Query/retrieval activity
        self.query_count = 0
        self.retrieval_nodes = 0
        self.query_duration = 0.0

    def attach(self):
        Settings.callback_manager = self.callback_manager

    def llm_snapshot(self):
        return (
            len(self.llm_token_counter.llm_token_counts),
            self.llm_token_counter.total_llm_token_count,
            self.llm_token_counter.prompt_llm_token_count,
            self.llm_token_counter.completion_llm_token_count,
        )

    def embedding_snapshot(self):
        handler = self.llm_token_counter
        return (
            len(getattr(handler, "embedding_token_counts", [])),
            getattr(handler, "total_embedding_token_count", 0),
        )

    def record_llm_operation(self, name, before_snapshot):
        after = self.llm_snapshot()
        calls = max(0, after[0] - before_snapshot[0])
        total = max(0, after[1] - before_snapshot[1])
        input_tokens = max(0, after[2] - before_snapshot[2])
        output_tokens = max(0, after[3] - before_snapshot[3])

        operation = self.llm_operations.setdefault(
            name,
            {"calls": 0, "input_tokens": 0, "output_tokens": 0, "total_tokens": 0},
        )
        operation["calls"] += calls
        operation["input_tokens"] += input_tokens
        operation["output_tokens"] += output_tokens
        operation["total_tokens"] += total
        self.llm_calls += calls

    def record_llama_parse(self, duration):
        self.llama_parse_calls += 1
        self.llama_parse_duration += duration

    def record_embedding_delta(self, before_snapshot, duration=0.0):
        after = self.embedding_snapshot()
        calls = max(0, after[0] - before_snapshot[0])
        tokens = max(0, after[1] - before_snapshot[1])
        self.embedding_calls += calls
        self.embedding_tokens += tokens
        self.embedding_duration += duration

    def record_indexing(self, duration, nodes):
        self.indexing_duration += duration
        self.nodes_indexed += nodes

    def record_chunking(self, duration):
        self.chunking_duration += duration

    def record_query(self, query_count, duration, retrieved_nodes):
        self.query_count += 1
        self.query_duration += duration
        self.retrieval_nodes += retrieved_nodes

    def query_usage_delta(self, before_snapshot):
        after = self.llm_snapshot()
        return (
            max(0, after[0] - before_snapshot[0]),
            max(0, after[2] - before_snapshot[2]),
            max(0, after[3] - before_snapshot[3]),
            max(0, after[1] - before_snapshot[1]),
        )

    def print_query_usage(self, before_snapshot):
        calls, input_tokens, output_tokens, total = self.query_usage_delta(before_snapshot)
        print("\n[QUERY OBSERVABILITY]")
        print(f"  Gemini API calls:   {calls}")
        print(f"  Input tokens:       {input_tokens:,}  (prompt + instructions + RAG context)")
        print(f"  Output tokens:      {output_tokens:,}  (Gemini-generated text)")
        print(f"  Total tokens:       {total:,}")

    def print_summary(self):
        input_tokens = self.llm_token_counter.prompt_llm_token_count
        output_tokens = self.llm_token_counter.completion_llm_token_count
        total_tokens = self.llm_token_counter.total_llm_token_count
        session_seconds = time.perf_counter() - self.started_at

        print("\n" + "=" * 70)
        print("                    PHASE 2 OBSERVABILITY")
        print("=" * 70)
        print(f"SESSION TIME:         {session_seconds:.2f}s")
        print("\nGEMINI")
        #print(f"  Model:              {GEMINI_MODEL}")
        print(f"  API calls:          {self.llm_calls}")
        print(f"  Input tokens:       {input_tokens:,}  (prompts + instructions + RAG context)")
        print(f"  Output tokens:      {output_tokens:,}  (Gemini-generated text)")
        print(f"  Total tokens:       {total_tokens:,}")
        print("  By operation:")
        for name, usage in self.llm_operations.items():
            print(
                f"    {name}: {usage['calls']} calls | "
                f"input {usage['input_tokens']:,} | "
                f"output {usage['output_tokens']:,} | "
                f"total {usage['total_tokens']:,}"
            )

        print("\nLLAMAPARSE")
        print(f"  API calls:          {self.llama_parse_calls}")
        print(f"  Parse time:         {self.llama_parse_duration:.2f}s")

        print("\nLOCAL PIPELINE")
        print(f"  Embedding calls:    {self.embedding_calls}")
        print(f"  Embedding tokens:   {self.embedding_tokens:,}")
        print(f"  Embedding time:     {self.embedding_duration:.2f}s")
        print(f"  Nodes indexed:      {self.nodes_indexed:,}")
        print(f"  Indexing time:      {self.indexing_duration:.2f}s")
        print(f"  Chunking time:      {self.chunking_duration:.2f}s")

        print("\nRETRIEVAL")
        print(f"  Queries:            {self.query_count}")
        print(f"  Retrieved nodes:    {self.retrieval_nodes:,}")
        print(f"  Query time:         {self.query_duration:.2f}s")
        print("=" * 70)


API_TRACKER = APIUsageTracker()
API_TRACKER.attach()


# ENVIRONMENT
load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
LLAMA_CLOUD_API_KEY = os.getenv("LLAMA_CLOUD_API_KEY")

if not GEMINI_API_KEY:
    raise ValueError(
        "GEMINI_API_KEY missing from .env file."
    )

if not LLAMA_CLOUD_API_KEY:
    raise ValueError(
        "LLAMA_CLOUD_API_KEY missing from .env file."
    )

# DIRECTORIES

DATA_DIR = Path("./data")

STORAGE_DIR = Path("./storage")

MANIFEST_FILE = STORAGE_DIR / "document_manifest.json"


DATA_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

STORAGE_DIR.mkdir(
    parents=True,
    exist_ok=True,
)

# MODELS

Settings.llm = GoogleGenAI(
    model="models/gemini-3.1-flash-lite"
)

Settings.embed_model = HuggingFaceEmbedding(
    model_name="BAAI/bge-small-en-v1.5"
)

# SUPPORTED FILE TYPES

LLAMA_PARSE_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".doc",
}

MARKDOWN_EXTENSIONS = {
    ".md",
    ".markdown",
}

CODE_LANGUAGES = {

    ".py": "python",

    ".js": "javascript",
    ".jsx": "javascript",

    ".ts": "typescript",
    ".tsx": "typescript",

    ".java": "java",

    ".cpp": "cpp",
    ".cc": "cpp",
    ".cxx": "cpp",

    ".c": "c",

    ".h": "c",
    ".hpp": "cpp",

    ".go": "go",

    ".rs": "rust",

    ".rb": "ruby",

    ".php": "php",

    ".swift": "swift",

    ".kt": "kotlin",
    ".kts": "kotlin",

    ".cs": "csharp",

    ".scala": "scala",

    ".sql": "sql",

    ".sh": "bash",
    ".bash": "bash",
}

TEXT_EXTENSIONS = {
    ".txt",
    ".text",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".env",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
}

JSON_EXTENSION = {
    ".json",
}

STRUCTURED_EXTENSIONS = {
    ".csv",
    ".tsv",
}

SPREADSHEET_EXTENSIONS = {
    ".xlsx",
    ".xls",
}

HTML_EXTENSIONS = {
    ".html",
    ".htm",
}

# DOCUMENT PROFILE

@dataclass
class DocumentProfile:

    path: Path

    extension: str

    file_size_bytes: int

    document_type: str = "unknown"

    word_count: int = 0

    line_count: int = 0

    has_headers: bool = False

    has_tables: bool = False

    has_code: bool = False

    structure_depth: int = 0

    language: Optional[str] = None

# HASH FUNCTION

def calculate_file_hash(
    path: Path,
) -> str:
    """
    Hash the actual file contents.

    Changes when the contents of the file change.
    """

    sha256 = hashlib.sha256()

    with open(
        path,
        "rb",
    ) as file:

        while True:

            chunk = file.read(
                1024 * 1024
            )

            if not chunk:
                break

            sha256.update(chunk)

    return sha256.hexdigest()


def calculate_document_id(
    relative_path: str,
) -> str:
    """
    Generate a stable document ID from the file path.

    IMPORTANT:

    This does NOT change when the file contents change.
    """

    normalized = relative_path.replace(
        "\\",
        "/",
    ).lower()

    return hashlib.sha256(
        normalized.encode("utf-8")
    ).hexdigest()

# MANIFEST

def load_manifest() -> Dict:

    if not MANIFEST_FILE.exists():
        return {}

    try:

        with open(
            MANIFEST_FILE,
            "r",
            encoding="utf-8",
        ) as file:

            return json.load(file)

    except Exception:

        return {}


def save_manifest(
    manifest: Dict,
):

    STORAGE_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    temp_file = MANIFEST_FILE.with_suffix(
        ".tmp"
    )

    try:
        
        with open(
            temp_file,
            "w",
            encoding="utf-8",
        ) as file:

            json.dump(
                manifest,
                file,
                indent=2,
                ensure_ascii=False,
            )

            file.flush()
            os.fsync(
                file.fileno()
            )

        os.replace(
            temp_file,
            MANIFEST_FILE,
        )

    except Exception:

        if temp_file.exists():
            temp_file.unlink()

        raise


# FILE CLASSIFICATION

def classify_file(
    path: Path,
) -> str:

    extension = path.suffix.lower()

    if extension in LLAMA_PARSE_EXTENSIONS:
        return "document"

    if extension in MARKDOWN_EXTENSIONS:
        return "markdown"

    if extension in CODE_LANGUAGES:
        return "code"

    if extension == ".json":
        return "json"

    if extension in {".csv", ".tsv"}:
        return "tabular"

    if extension in SPREADSHEET_EXTENSIONS:
        return "spreadsheet"

    if extension in HTML_EXTENSIONS:
        return "html"

    if extension in TEXT_EXTENSIONS:
        return "text"

    return "unknown"

# TEXT ANALYSIS

def run_local_metrics(text: str) -> Tuple[int, int]:
    
    words = re.findall(r"\b\w+\b", text)
    lines = text.splitlines()
    return len(words), len(lines)


def run_heuristic_filter(text: str) -> dict:

    has_code_signals = bool(re.search(
        r"\b(def|class|import|function|SELECT|FROM)\b|=>|if\s*\(", 
        text, 
        flags=re.IGNORECASE
    ))
    
    has_table_signals = bool(re.search(
        r"\||(?m)^[^,\n]*(?:,[^,\n]*){3,}|(?m)^[^\t\n]*(?:\t[^\t\n]*){3,}", 
        text
    ))
    
    has_header_signals = bool(re.search(
        r"(?m)^#{1,6}\s+|(?m)^\d+(\.\d+)*\s+[A-Z]", 
        text
    ))

    return {
        "needs_code_check": has_code_signals,
        "needs_table_check": has_table_signals,
        "needs_header_check": has_header_signals
    }


def evaluate_content_structure_llm(text: str, filters: dict) -> dict:
   
    structure = {
        "has_headers": False,
        "has_tables": False,
        "has_code": False
    }

    checks_to_perform = []
    if filters["needs_header_check"]:
        checks_to_perform.append("1. Structural headings, section titles, or chapter lines.")
    if filters["needs_table_check"]:
        checks_to_perform.append("2. Data grids, tables, CSV formatting, or spreadsheet layouts.")
    if filters["needs_code_check"]:
        checks_to_perform.append("3. Functional blocks of source code, configurations, or scripting syntax.")

    if not checks_to_perform or not text.strip():
        return structure

    snippet = text[:2000]
    checks_str = "\n".join(checks_to_perform)
    
    prompt = f"""
    Analyze the following text snippet from a document. Determine if it contains any of these specific elements:
    {checks_str}

    Respond STRICTLY in valid raw JSON format matching this schema without any markdown formatting wrappers:
    {{
        "has_headers": true/false,
        "has_tables": true/false,
        "has_code": true/false
    }}

    Text Snippet:
    \"\"\"{snippet}\"\"\"
    """

    try:
        llm_before = API_TRACKER.llm_snapshot()
        response = Settings.llm.complete(prompt)
        API_TRACKER.record_llm_operation(
            "document_profiling",
            llm_before,
        )
        clean_response = repair_json(response.text.strip())
        
        
        llm_analysis = json.loads(clean_response)
        structure.update(llm_analysis)
    except Exception:
        pass

    return structure


def profile_document(
    path: Path,
    extracted_text: Optional[str] = None,
) -> DocumentProfile:
    
    extension = path.suffix.lower()
    document_type = classify_file(path)

    profile = DocumentProfile(
        path=path,
        extension=extension,
        file_size_bytes=path.stat().st_size,
        document_type=document_type,
    )

    if extension in CODE_LANGUAGES:
        profile.language = CODE_LANGUAGES[extension]
        profile.has_code = True

    if extracted_text:
        profile.word_count, profile.line_count = run_local_metrics(extracted_text)

        filters = run_heuristic_filter(extracted_text)
        
        if filters["needs_code_check"] or filters["needs_table_check"] or filters["needs_header_check"]:
            llm_insights = evaluate_content_structure_llm(extracted_text, filters)
            
            profile.has_headers = llm_insights["has_headers"]
            profile.has_tables = llm_insights["has_tables"]
            if not profile.has_code:
                profile.has_code = llm_insights["has_code"]
        else:
            profile.has_headers = False
            profile.has_tables = False

        header_matches = re.findall(r"(?m)^#{1,6}\s+", extracted_text)
        if header_matches:
            profile.structure_depth = max(len(header.strip()) for header in header_matches)
        elif profile.has_headers:
            profile.structure_depth = 1

    return profile

# LLAMAPARSE

def create_llama_parser():

    return LlamaParse(
        result_type="markdown",
        verbose=False,
    )

# DOCUMENT LOADING

def attach_document_identity(
    document: Document,
    path: Path,
    document_id: str,
    file_hash: str,
):

    document.doc_id = document_id

    document.metadata.update(
        {
            "document_id": document_id,
            "file_hash": file_hash,
            "file_name": path.name,
            "file_path": str(path),
        }
    )


def load_single_file(
    path: Path,
    document_id: str,
    file_hash: str,
) -> List[Document]:

    extension = path.suffix.lower()

    document_type = classify_file(
        path
    )

    if extension in LLAMA_PARSE_EXTENSIONS:

        parser = create_llama_parser()

        parse_start = time.perf_counter()
        parsed_documents = parser.load_data(
            str(path)
        )
        API_TRACKER.record_llama_parse(
            time.perf_counter() - parse_start
        )

        for document in parsed_documents:

            attach_document_identity(
                document,
                path,
                document_id,
                file_hash,
            )

            document.metadata[
                "document_type"
            ] = document_type

        return parsed_documents

    if extension in MARKDOWN_EXTENSIONS:

        text = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

        document = Document(
            text=text,
        )

        attach_document_identity(
            document,
            path,
            document_id,
            file_hash,
        )

        document.metadata[
            "document_type"
        ] = document_type

        return [document]

    if extension in CODE_LANGUAGES:

        text = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

        document = Document(
            text=text,
        )

        attach_document_identity(
            document,
            path,
            document_id,
            file_hash,
        )

        document.metadata.update(
            {
                "document_type": document_type,
                "language": CODE_LANGUAGES[
                    extension
                ],
            }
        )

        return [document]

    if extension == ".json":

        try:

            with open(
                path,
                "r",
                encoding="utf-8",
            ) as file:

                data = json.load(file)

            text = json.dumps(
                data,
                indent=2,
                ensure_ascii=False,
            )

        except Exception:

            text = path.read_text(
                encoding="utf-8",
                errors="ignore",
            )

        document = Document(
            text=text,
        )

        attach_document_identity(
            document,
            path,
            document_id,
            file_hash,
        )

        document.metadata[
            "document_type"
        ] = document_type

        return [document]

    if extension in {
        ".csv",
        ".tsv",
    }:

        delimiter = (
            "\t"
            if extension == ".tsv"
            else ","
        )

        rows = []

        try:

            with open(
                path,
                "r",
                encoding="utf-8",
                errors="ignore",
            ) as file:

                reader = csv.reader(
                    file,
                    delimiter=delimiter,
                )

                for row in reader:

                    rows.append(
                        " | ".join(
                            str(value)
                            for value in row
                        )
                    )

        except Exception:

            rows = [
                path.read_text(
                    encoding="utf-8",
                    errors="ignore",
                )
            ]

        document = Document(
            text="\n".join(rows)
        )

        attach_document_identity(
            document,
            path,
            document_id,
            file_hash,
        )

        document.metadata[
            "document_type"
        ] = document_type

        return [document]

    if extension in SPREADSHEET_EXTENSIONS:

        return load_excel_file(
            path,
            document_id,
            file_hash,
        )

    if extension in HTML_EXTENSIONS:
        
        reader = HTMLTagReader()
        parsed_documents = reader.load_data(file=path)
        
        for document in parsed_documents:
            attach_document_identity(
                document,
                path, 
                document_id, 
                file_hash
            )
            document.metadata[
                "document_type"
            ] = document_type

        return parsed_documents

    try:

        text = path.read_text(
            encoding="utf-8",
            errors="ignore",
        )

        if not text.strip():
            return []

        document = Document(
            text=text,
        )

        attach_document_identity(
            document,
            path,
            document_id,
            file_hash,
        )

        document.metadata[
            "document_type"
        ] = document_type

        return [document]

    except Exception as error:

        print(
            f"--> [SKIP] {path.name}: {error}"
        )

        return []


def load_excel_file(
    path: Path,
    document_id: str,
    file_hash: str,
) -> List[Document]:

    try:

        from openpyxl import load_workbook

    except ImportError:

        raise ImportError(
            "Install openpyxl to process Excel files."
        )

    documents = []

    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=True,
    )

    for worksheet in workbook.worksheets:

        rows = []

        for row in worksheet.iter_rows(
            values_only=True
        ):

            values = [
                ""
                if value is None
                else str(value)
                for value in row
            ]

            if any(values):

                rows.append(
                    " | ".join(values)
                )

        text = "\n".join(rows)

        if not text.strip():
            continue

        document = Document(
            text=text
        )

        attach_document_identity(
            document,
            path,
            document_id,
            file_hash,
        )

        document.metadata.update(
            {
                "document_type": "spreadsheet",
                "sheet_name": worksheet.title,
            }
        )

        documents.append(
            document
        )

    return documents


class AdaptiveChunker:

    def process(
        self,
        documents: List[Document],
    ):

        nodes = []

        for document in documents:

            path = Path(
                document.metadata.get(
                    "file_path",
                    "unknown",
                )
            )

            profile = profile_document(
                path,
                extracted_text=document.text,
            )

            document.metadata.update(
                {
                    "word_count": profile.word_count,
                    "line_count": profile.line_count,
                    "has_headers": profile.has_headers,
                    "has_tables": profile.has_tables,
                    "has_code": profile.has_code,
                    "structure_depth": profile.structure_depth,
                }
            )

            strategy = self.choose_strategy(
                profile
            )

            print(
                f"--> [PROFILE] "
                f"{path.name} | "
                f"type={profile.document_type} | "
                f"words={profile.word_count} | "
                f"strategy={strategy}"
            )

            generated_nodes = (
                self.chunk_document(
                    document,
                    profile,
                    strategy,
                )
            )

            nodes.extend(
                generated_nodes
            )

        return nodes

    def choose_strategy(
        self,
        profile: DocumentProfile,
    ) -> str:

        if profile.document_type == "code":
            return "code"

        if profile.document_type == "markdown":
            return "markdown"

        if profile.document_type == "document":

            if (
                profile.has_headers
                or profile.has_tables
            ):
                return "structured_document"

            return "document"

        if profile.document_type in {
            "json",
            "tabular",
            "spreadsheet",
        }:
            return "structured_data"

        if profile.document_type == "html":
            return "html"

        if profile.word_count < 700:
            return "short_text"

        return "adaptive_text"

    def chunk_document(
        self,
        document: Document,
        profile: DocumentProfile,
        strategy: str,
    ):

        if strategy == "code":

            return self.chunk_code(
                document,
                profile,
            )

        if strategy in {
            "markdown",
            "structured_document",
        }:

            return self.chunk_markdown(
                document
            )

        if strategy == "structured_data":

            return self.chunk_structured(
                document
            )

        if strategy == "short_text":

            return [
                self.set_strategy(
                    document,
                    "minimal",
                )
            ]

        return self.chunk_text(
            document,
            profile,
        )

    def chunk_code(
        self,
        document: Document,
        profile: DocumentProfile,
    ):

        language = (
            profile.language
            or "python"
        )

        try:

            splitter = CodeSplitter(
                language=language,
                chunk_lines=self.dynamic_code_lines(
                    profile.line_count
                ),
            )

            pipeline = IngestionPipeline(
                transformations=[
                    splitter
                ]
            )

            nodes = pipeline.run(
                documents=[document]
            )

            for node in nodes:

                node.metadata.update(
                    {
                        "chunk_strategy": "code",
                        "language": language,
                    }
                )

            return nodes

        except Exception as error:

            print(
                f"--> [WARNING] "
                f"Code splitter failed: {error}"
            )

            return self.chunk_text(
                document,
                profile,
            )

    def dynamic_code_lines(
        self,
        line_count: int,
    ) -> int:

        if line_count < 100:
            return 80

        if line_count < 500:
            return 60

        if line_count < 2000:
            return 50

        if line_count < 5000:
            return 40

        return 30

    def chunk_markdown(
        self,
        document: Document,
    ):

        pipeline = IngestionPipeline(
            transformations=[
                MarkdownNodeParser()
            ]
        )

        nodes = pipeline.run(
            documents=[document]
        )

        for node in nodes:

            node.metadata[
                "chunk_strategy"
            ] = "markdown_structure"

        return nodes

    def chunk_structured(
        self,
        document: Document,
    ):

        word_count = len(
            re.findall(
                r"\b\w+\b",
                document.text,
            )
        )

        if word_count < 2000:

            return [
                self.set_strategy(
                    document,
                    "structured_data",
                )
            ]

        return self.chunk_text(
            document,
            DocumentProfile(
                path=Path(
                    document.metadata[
                        "file_path"
                    ]
                ),
                extension=document.metadata.get(
                    "extension",
                    "",
                ),
                file_size_bytes=0,
                document_type=document.metadata.get(
                    "document_type", 
                    "tabular"
                ),
                word_count=word_count,
            ),
        )

    def chunk_text(
        self,
        document: Document,
        profile: DocumentProfile,
    ):

        chunk_size = (
            self.dynamic_chunk_size(
                profile.word_count
            )
        )

        overlap = (
            self.dynamic_overlap(
                chunk_size
            )
        )

        splitter = SentenceSplitter(
            chunk_size=chunk_size,
            chunk_overlap=overlap,
        )

        pipeline = IngestionPipeline(
            transformations=[
                splitter
            ]
        )

        nodes = pipeline.run(
            documents=[document]
        )

        for node in nodes:

            node.metadata[
                "chunk_strategy"
            ] = "adaptive_text"

        return nodes

    def dynamic_chunk_size(
        self,
        word_count: int,
    ) -> int:

        if word_count < 1000:
            return 700

        if word_count < 5000:
            return 650

        if word_count < 20000:
            return 550

        if word_count < 100000:
            return 450

        return 350

    def dynamic_overlap(
        self,
        chunk_size: int,
    ) -> int:

        return max(
            30,
            int(chunk_size * 0.10),
        )

    def set_strategy(
        self,
        document,
        strategy,
    ):

        document.metadata[
            "chunk_strategy"
        ] = strategy

        return document
    

def discover_files(
    data_dir: Path,
) -> List[Path]:

    files = []

    for path in data_dir.rglob("*"):

        if not path.is_file():
            continue

        if path.name.startswith("."):
            continue

        if STORAGE_DIR in path.parents:
            continue

        files.append(path)

    return sorted(files)



@dataclass
class QueryProfile:
    """Lightweight query analysis used to select an adaptive retrieval path."""

    mode: str
    reason: str
    needs_exact_match: bool = False
    is_broad: bool = False
    is_complex: bool = False


class LexicalIndex:
    """
    Small dependency-free BM25-style lexical index.

    It indexes node text plus useful metadata so exact identifiers, filenames,
    fields, error codes, SKUs, and similar lexical signals can be retrieved
    even when dense similarity is not the strongest signal.
    """

    def __init__(self, k1: float = 1.5, b: float = 0.75):
        self.k1 = k1
        self.b = b
        self.nodes = []
        self.term_frequencies = []
        self.document_frequencies = Counter()
        self.document_lengths = []
        self.average_document_length = 0.0
        self.built = False

    @staticmethod
    def tokenize(text: str) -> List[str]:
        return re.findall(r"[A-Za-z0-9_]+", text.lower())

    @staticmethod
    def searchable_text(node) -> str:
        metadata = getattr(node, "metadata", {}) or {}
        metadata_parts = [
            str(metadata.get(key, ""))
            for key in (
                "file_name",
                "file_path",
                "sheet_name",
                "language",
                "document_type",
                "chunk_strategy",
            )
        ]
        return " ".join(
            [node.get_content()] + metadata_parts
        )

    def build(self, nodes: List) -> None:
        self.nodes = list(nodes)
        self.term_frequencies = []
        self.document_frequencies = Counter()
        self.document_lengths = []

        for node in self.nodes:
            terms = self.tokenize(self.searchable_text(node))
            frequencies = Counter(terms)
            self.term_frequencies.append(frequencies)
            self.document_lengths.append(len(terms))
            self.document_frequencies.update(frequencies.keys())

        total_length = sum(self.document_lengths)
        self.average_document_length = (
            total_length / len(self.nodes)
            if self.nodes
            else 0.0
        )
        self.built = True

    def retrieve(self, query: str, top_k: int = 8) -> List[NodeWithScore]:
        if not self.built or not self.nodes:
            return []

        query_terms = self.tokenize(query)
        if not query_terms:
            return []

        query_counter = Counter(query_terms)
        total_documents = len(self.nodes)
        average_length = self.average_document_length or 1.0

        scored = []
        for index, node in enumerate(self.nodes):
            frequencies = self.term_frequencies[index]
            document_length = self.document_lengths[index] or 1
            score = 0.0

            for term in query_counter:
                frequency = frequencies.get(term, 0)
                if frequency == 0:
                    continue

                document_frequency = self.document_frequencies.get(term, 0)
                idf = math.log(
                    1.0
                    + (
                        (total_documents - document_frequency + 0.5)
                        / (document_frequency + 0.5)
                    )
                )

                denominator = (
                    frequency
                    + self.k1
                    * (
                        1.0
                        - self.b
                        + self.b
                        * document_length
                        / average_length
                    )
                )
                score += (
                    idf
                    * (
                        frequency
                        * (self.k1 + 1.0)
                        / denominator
                    )
                )

            if score <= 0:
                continue

            searchable = self.searchable_text(node).lower()
            query_text = query.lower().strip()
            if query_text and query_text in searchable:
                score += 1.0

            scored.append(
                NodeWithScore(
                    node=node,
                    score=float(score),
                )
            )

        scored.sort(
            key=lambda item: item.score or 0.0,
            reverse=True,
        )
        return scored[:top_k]


class AdaptiveRAG:

    VECTOR_INDEX_ID = "vector_idx"

    SUMMARY_INDEX_ID = "summary_idx"

    def __init__(
        self,
        data_dir: str = "./data",
    ):

        self.data_dir = Path(
            data_dir
        )

        self.manifest = (
            load_manifest()
        )

        self.chunker = (
            AdaptiveChunker()
        )

        self.engine = None

        self.vector_index = None

        self.summary_index = None

        self.db_client = qdrant_client.QdrantClient(url="http://localhost:6333")

        # Phase 5: adaptive retrieval configuration.
        self.lexical_index = LexicalIndex()
        self.retrieval_top_k = 5
        self.candidate_top_k = 12
        self.default_rag_mode = "auto"

    def indexes_exist(self) -> bool:

        try:
            collections = self.db_client.get_collections().collections
            exists = any(c.name == "pipeline_collection" for c in collections)
            return exists and (STORAGE_DIR / "docstore.json").exists()
        except Exception:
            return False

    def load_indexes(self):

        if not self.indexes_exist():
            return False

        print(
            "--> [LOAD] Connecting to existing Local VectorDB..."
        )

        vector_store = QdrantVectorStore(
            client=self.db_client, 
            collection_name="pipeline_collection"
        )

        storage_context = (
            StorageContext.from_defaults(
                vector_store=vector_store,
                persist_dir=str(
                    STORAGE_DIR
                )
            )
        )

        self.vector_index = (
            load_index_from_storage(
                storage_context,
                index_id=self.VECTOR_INDEX_ID,
            )
        )

        self.summary_index = (
            load_index_from_storage(
                storage_context,
                index_id=self.SUMMARY_INDEX_ID,
            )
        )

        self.rebuild_lexical_index()

        return True

    def persist_indexes(self):

        if self.vector_index is None and self.summary_index is None:
            return

        print("--> [PERSIST] Saving index state...")

        self.vector_index.storage_context.persist(
            persist_dir=str(STORAGE_DIR)
        )

        self.summary_index.storage_context.persist(
            persist_dir=str(STORAGE_DIR)
        )

        print("--> [PERSIST] Index state saved.")

    def create_indexes(
        self,
        nodes,
    ):

        print(
            "--> [INDEX] Creating production VectorDB collection..."
        )

        index_start = time.perf_counter()
        embedding_before = API_TRACKER.embedding_snapshot()

        vector_store = QdrantVectorStore(
            client=self.db_client,
            collection_name="pipeline_collection"
        )

        storage_context = (
            StorageContext.from_defaults(vector_store=vector_store)
        )

        self.vector_index = (
            VectorStoreIndex(
                nodes,
                storage_context=storage_context,
            )
        )

        self.vector_index.set_index_id(
            self.VECTOR_INDEX_ID
        )

        self.summary_index = (
            SummaryIndex(
                nodes,
                storage_context=storage_context,
            )
        )

        self.summary_index.set_index_id(
            self.SUMMARY_INDEX_ID
        )

        self.persist_indexes()

        API_TRACKER.record_indexing(
            time.perf_counter() - index_start,
            len(nodes),
        )
        API_TRACKER.record_embedding_delta(embedding_before)
        self.rebuild_lexical_index()

    def delete_document_from_indexes(
        self,
        document_id: str,
    ):

        print(
            f"--> [DELETE INDEX] "
            f"Removing document "
            f"{document_id[:12]}..."
        )

        if self.vector_index is None:
            raise RuntimeError(
            "Vector index is not loaded."
        )

        try:

            self.vector_index.delete_ref_doc(
                ref_doc_id=document_id,
                delete_from_docstore=True,
            )

            print(
                "--> [DELETE INDEX] "
                "Vector index cleaned."
            )

        except Exception as error:
            raise RuntimeError(
                f"Vector index deletion failed for "
                f"{document_id}: {error}"
            ) from error


        if self.summary_index:

            try:

                self.summary_index.delete_ref_doc(
                    ref_doc_id=document_id,
                    delete_from_docstore=True,
                )

                print(
                    "--> [DELETE INDEX] "
                    "Summary index cleaned."
                )

            except Exception as error:
                raise RuntimeError(
                    f"Summary index deletion failed for "
                    f"{document_id}: {error}"
                ) from error
            
        self.persist_indexes()
        self.rebuild_lexical_index()


    def ingest_file(
        self,
        path: Path,
        relative_path: str,
        document_id: str,
        file_hash: str,
    ):

        print(
            f"--> [INGEST] {relative_path}"
        )

        documents = load_single_file(
            path=path,
            document_id=document_id,
            file_hash=file_hash,
        )

        if not documents:

            print(
                "--> [WARNING] "
                f"No readable content: "
                f"{relative_path}"
            )

            return []

        chunk_start = time.perf_counter()
        nodes = self.chunker.process(
            documents
        )

        for node in nodes:
            node.id_ = str(uuid.uuid4())
            node.relationships[NodeRelationship.SOURCE] = RelatedNodeInfo(node_id=document_id)

        chunk_seconds = time.perf_counter() - chunk_start
        API_TRACKER.record_chunking(chunk_seconds)

        print(
            f"--> [INGEST] "
            f"{relative_path} → "
            f"{len(nodes)} nodes"
        )

        return nodes

    def sync(
        self,
        force_rebuild: bool = False,
    ):

        print()
        print("=" * 70)
        print("              ADAPTIVE RAG SYNC")
        print("=" * 70)

        if force_rebuild:

            print(
                "--> [REBUILD] "
                "Complete rebuild requested."
            )

            self.clear_storage()

            self.manifest = {}

            self.vector_index = None

            self.summary_index = None

        else:

            self.load_indexes()


        files = discover_files(
            self.data_dir
        )

        current_files = {}

        for path in files:

            relative_path = str(
                path.relative_to(
                    self.data_dir
                )
            )

            document_id = (
                calculate_document_id(
                    relative_path
                )
            )

            file_hash = (
                calculate_file_hash(
                    path
                )
            )

            current_files[
                relative_path
            ] = {
                "document_id": document_id,
                "hash": file_hash,
                "extension": path.suffix.lower(),
                "size": path.stat().st_size,
            }

        if not self.indexes_exist():

            print(
                "--> [INIT] "
                "No indexes found. "
                "Building initial RAG."
            )

            all_nodes = []

            for path in files:

                relative_path = str(
                    path.relative_to(
                        self.data_dir
                    )
                )

                info = current_files[
                    relative_path
                ]

                nodes = self.ingest_file(
                    path=path,
                    relative_path=relative_path,
                    document_id=info[
                        "document_id"
                    ],
                    file_hash=info[
                        "hash"
                    ],
                )

                all_nodes.extend(
                    nodes
                )

            if all_nodes:

                self.create_indexes(
                    all_nodes
                )

            self.manifest = (
                current_files
            )

            save_manifest(
                self.manifest
            )

            print(
                "--> [DONE] "
                "Initial RAG created."
            )

            self.build_router()

            return

        old_paths = set(
            self.manifest.keys()
        )

        current_paths = set(
            current_files.keys()
        )

        deleted_paths = (
            old_paths - current_paths
        )

        for relative_path in deleted_paths:

            old_record = (
                self.manifest[
                    relative_path
                ]
            )

            document_id = (
                old_record[
                    "document_id"
                ]
            )

            print(
                f"--> [DELETED FILE] "
                f"{relative_path}"
            )

            self.delete_document_from_indexes(
                document_id
            )

            del self.manifest[
                relative_path
            ]

        for relative_path, info in (
            current_files.items()
        ):

            old_record = (
                self.manifest.get(
                    relative_path
                )
            )

            path = (
                self.data_dir
                / relative_path
            )

            if old_record is None:

                print(
                    f"--> [NEW FILE] "
                    f"{relative_path}"
                )

                nodes = self.ingest_file(
                    path=path,
                    relative_path=relative_path,
                    document_id=info[
                        "document_id"
                    ],
                    file_hash=info[
                        "hash"
                    ],
                )

                self.insert_nodes(
                    nodes
                )

                self.manifest[
                    relative_path
                ] = info

                continue

            old_hash = (
                old_record.get(
                    "hash"
                )
            )

            new_hash = (
                info["hash"]
            )

            if old_hash != new_hash:

                print(
                    f"--> [CHANGED FILE] "
                    f"{relative_path}"
                )

                document_id = (
                    old_record[
                        "document_id"
                    ]
                )

                self.delete_document_from_indexes(
                    document_id
                )

                nodes = self.ingest_file(
                    path=path,
                    relative_path=relative_path,
                    document_id=document_id,
                    file_hash=new_hash,
                )

                self.insert_nodes(
                    nodes
                )

                self.manifest[
                    relative_path
                ] = {
                    **info,
                    "document_id": document_id,
                }

                continue

            print(
                f"--> [UNCHANGED] "
                f"{relative_path}"
            )

        save_manifest(
            self.manifest
        )

        print()
        print(
            "--> [DONE] "
            "RAG synchronized successfully."
        )

        self.build_router()

    def insert_nodes(
        self,
        nodes,
    ):

        if not nodes:
            return

        if self.vector_index is None:

            print(
                "--> [INDEX] "
                "Vector index doesn't exist."
            )

            self.create_indexes(
                nodes
            )

            return

        print(
            f"--> [INDEX] "
            f"Inserting {len(nodes)} nodes..."
        )

        try:

            index_start = time.perf_counter()
            embedding_before = API_TRACKER.embedding_snapshot()

            self.vector_index.insert_nodes(
                nodes
            )

            if self.summary_index is not None:
                self.summary_index.insert_nodes(
                    nodes
                )

            self.persist_indexes()
            API_TRACKER.record_indexing(
            time.perf_counter() - index_start,
            len(nodes),
            )
            API_TRACKER.record_embedding_delta(embedding_before)
            self.rebuild_lexical_index()

        except Exception as error:

            raise RuntimeError(
                f"Failed to insert nodes into indexes: "
                f"{error}"
            ) from error

    def rebuild_lexical_index(self):
        """Rebuild the local lexical index from the persisted node docstore."""
        if self.vector_index is None:
            self.lexical_index = LexicalIndex()
            return

        nodes = list(self.vector_index.docstore.docs.values())
        self.lexical_index.build(nodes)

        print(
            f"--> [RETRIEVAL] Lexical index rebuilt: "
            f"{len(nodes)} nodes"
        )

    def profile_query(
        self,
        question: str,
        rag_mode: str = "auto",
    ) -> QueryProfile:
        """
        Select a retrieval strategy.

        User-selected modes are honored explicitly. In auto mode the system
        uses lightweight lexical signals to choose between semantic, keyword,
        hybrid, and summary retrieval.
        """
        requested_mode = (rag_mode or self.default_rag_mode).strip().lower()

        aliases = {
            "vector": "semantic",
            "dense": "semantic",
            "bm25": "keyword",
            "lexical": "keyword",
            "hybrid_search": "hybrid",
            "global": "summary",
        }
        requested_mode = aliases.get(
            requested_mode,
            requested_mode,
        )

        valid_modes = {
            "auto",
            "semantic",
            "keyword",
            "hybrid",
            "summary",
        }
        if requested_mode not in valid_modes:
            raise ValueError(
                f"Unsupported RAG mode '{rag_mode}'. "
                f"Choose from: {', '.join(sorted(valid_modes))}."
            )

        if requested_mode != "auto":
            return QueryProfile(
                mode=requested_mode,
                reason="user selected",
            )

        normalized = question.lower().strip()

        exact_patterns = [
            r"`[^`]+`",
            r"\b(error|exception|traceback|id|code|filename|"
            r"function|method|variable|class|field|column|sku|part)\b",
            r"\b[a-zA-Z][a-zA-Z0-9]*_[a-zA-Z0-9_]+\b",
            r"\b[a-zA-Z0-9]+-[a-zA-Z0-9-]+\b",
        ]
        needs_exact_match = any(
            re.search(pattern, question, flags=re.IGNORECASE)
            for pattern in exact_patterns
        ) or len(re.findall(r"\b\d+\b", question)) >= 2

        broad_terms = {
            "summarize",
            "summary",
            "overview",
            "overall",
            "main themes",
            "key themes",
            "entire document",
            "whole document",
            "all documents",
            "big picture",
        }
        is_broad = any(term in normalized for term in broad_terms)

        complex_terms = {
            "compare",
            "comparison",
            "contrast",
            "difference",
            "differences",
            "versus",
            "vs",
            "across",
            "multiple",
            "both",
            "and explain",
            "why and how",
        }
        is_complex = any(
            term in normalized
            for term in complex_terms
        ) or question.count("?") > 1

        if is_broad:
            return QueryProfile(
                mode="summary",
                reason="broad/global question",
                needs_exact_match=needs_exact_match,
                is_broad=True,
                is_complex=is_complex,
            )

        if needs_exact_match and is_complex:
            return QueryProfile(
                mode="hybrid",
                reason="exact-match signals + complex comparison",
                needs_exact_match=True,
                is_complex=True,
            )

        if needs_exact_match:
            return QueryProfile(
                mode="keyword",
                reason="exact identifier/fact signals",
                needs_exact_match=True,
            )

        if is_complex:
            return QueryProfile(
                mode="hybrid",
                reason="multi-part/comparison question",
                is_complex=True,
            )

        return QueryProfile(
            mode="semantic",
            reason="conceptual/semantic question",
        )

    def _vector_retrieve(
        self,
        question: str,
        top_k: int,
    ) -> List[NodeWithScore]:
        retriever = self.vector_index.as_retriever(
            similarity_top_k=top_k
        )
        return retriever.retrieve(question)

    def _rerank(
        self,
        question: str,
        candidates: List[NodeWithScore],
        top_k: int,
    ) -> List[NodeWithScore]:
        """
        Lightweight local reranker.

        It preserves the retriever score while rewarding query-term coverage
        and exact phrase matches. This keeps Phase 5 dependency-free; a
        cross-encoder can be evaluated later if testing shows it is needed.
        """
        query_terms = set(
            LexicalIndex.tokenize(question)
        )
        query_phrase = question.lower().strip()

        if not candidates:
            return []

        source_scores = [
            float(candidate.score or 0.0)
            for candidate in candidates
        ]
        max_score = max(source_scores) or 1.0

        reranked = []
        for candidate in candidates:
            content = candidate.node.get_content().lower()
            candidate_terms = set(
                LexicalIndex.tokenize(content)
            )

            coverage = (
                len(query_terms & candidate_terms)
                / len(query_terms)
                if query_terms
                else 0.0
            )
            exact_phrase = (
                1.0
                if query_phrase and query_phrase in content
                else 0.0
            )
            normalized_source = (
                float(candidate.score or 0.0) / max_score
            )

            final_score = (
                0.55 * normalized_source
                + 0.35 * coverage
                + 0.10 * exact_phrase
            )

            reranked.append(
                NodeWithScore(
                    node=candidate.node,
                    score=final_score,
                )
            )

        reranked.sort(
            key=lambda item: item.score or 0.0,
            reverse=True,
        )
        return reranked[:top_k]

    def _merge_hybrid(
        self,
        vector_results: List[NodeWithScore],
        keyword_results: List[NodeWithScore],
        top_k: int,
    ) -> List[NodeWithScore]:
        """Fuse dense and lexical candidates using reciprocal rank fusion."""
        fused_scores = defaultdict(float)
        nodes_by_id = {}

        for results in (vector_results, keyword_results):
            for rank, result in enumerate(results, start=1):
                node_id = result.node.node_id
                nodes_by_id[node_id] = result.node
                fused_scores[node_id] += 1.0 / (60.0 + rank)

        fused = [
            NodeWithScore(
                node=nodes_by_id[node_id],
                score=score,
            )
            for node_id, score in fused_scores.items()
        ]
        fused.sort(
            key=lambda item: item.score or 0.0,
            reverse=True,
        )
        return fused[:top_k]

    def retrieve_adaptively(
        self,
        question: str,
        profile: QueryProfile,
    ):
        """Run only the retrieval path selected for the current query."""
        if profile.mode == "summary":
            return None, "summary"

        if profile.mode == "semantic":
            candidates = self._vector_retrieve(
                question,
                self.candidate_top_k,
            )

        elif profile.mode == "keyword":
            candidates = self.lexical_index.retrieve(
                question,
                self.candidate_top_k,
            )

        elif profile.mode == "hybrid":
            vector_results = self._vector_retrieve(
                question,
                self.candidate_top_k,
            )
            keyword_results = self.lexical_index.retrieve(
                question,
                self.candidate_top_k,
            )
            candidates = self._merge_hybrid(
                vector_results,
                keyword_results,
                self.candidate_top_k,
            )

        else:
            raise ValueError(
                f"Unsupported retrieval mode: {profile.mode}"
            )

        reranked = self._rerank(
            question,
            candidates,
            self.retrieval_top_k,
        )

        # Expand each winning chunk by its immediate neighboring chunks when
        # the parser supplied PREVIOUS/NEXT relationships.
        postprocessor = PrevNextNodePostprocessor(
            docstore=self.vector_index.docstore,
            num_nodes=1,
            mode="both",
        )
        expanded = postprocessor.postprocess_nodes(reranked)

        # Deduplicate while preserving reranked/expanded order.
        unique_nodes = []
        seen_ids = set()
        for node_with_score in expanded:
            node_id = node_with_score.node.node_id
            if node_id in seen_ids:
                continue
            seen_ids.add(node_id)
            unique_nodes.append(node_with_score)

        return unique_nodes, profile.mode

    def synthesize_answer(
        self,
        question: str,
        retrieved_nodes: List[NodeWithScore],
    ) -> str:
        """Generate an answer strictly from the selected retrieval context."""
        if not retrieved_nodes:
            return (
                "The provided documentation does not contain "
                "enough information to answer this question."
            )

        contexts = []
        for index, item in enumerate(retrieved_nodes, start=1):
            node = item.node
            metadata = getattr(node, "metadata", {}) or {}
            source = (
                metadata.get("file_name")
                or metadata.get("file_path")
                or "unknown source"
            )
            contexts.append(
                f"[Context {index} | Source: {source}]\n"
                f"{node.get_content()}"
            )

        unified_context = "\n\n---\n\n".join(contexts)

        prompt = f"""
            You are a retrieval-augmented assistant.

            Answer the user's question using ONLY the supplied CONTEXT.
            Do not use outside knowledge.
            Do not invent facts, values, identifiers, filenames, or relationships.
            If the context does not contain enough information, say:
            "The provided documentation does not contain this information."

            QUESTION:
            {question}

            CONTEXT:
            \"\"\"{unified_context}\"\"\"

            Provide a concise, factual answer.
            """

        return Settings.llm.complete(prompt).text.strip()

    def build_router(self):
        """
        Prepare retrieval state.

        Phase 5 no longer asks an LLM to choose only between vector and
        summary engines. Query profiling now selects semantic, keyword,
        hybrid, or summary retrieval explicitly or automatically.
        """
        if self.vector_index is None:
            if not self.load_indexes():
                return None

        self.rebuild_lexical_index()

        if self.summary_index is not None:
            self.engine = self.summary_index.as_query_engine(
                response_mode="tree_summarize",
            )
        else:
            self.engine = None

        return self.engine

    def evaluate_faithfulness(self, answer: str, contexts: list) -> bool:
        
        if not contexts:
            return False
        
        unified_context = "\n---\n".join(contexts)
        
        eval_prompt = f"""
        You are a strict, independent Quality Assurance Auditor.
        Your sole task is to compare the given ANSWER against the verified CONTEXT text blocks.
        Determine if the ANSWER contains any fabrications, assumptions, extrapolations, or claims NOT explicitly proven by the CONTEXT.

        CONTEXT:
        \"\"\"{unified_context}\"\"\"

        ANSWER:
        \"\"\"{answer}\"\"\"

        Respond STRICTLY with one single word: 
        - Respond 'SAFE' if every single claim in the answer is 100% accurate and proven by the context.
        - Respond 'HALLUCINATION' if the answer contains any unproven claims, guesses, or inferences.
        
        Do not include any introductory remarks, punctuation, explanations, or markdown boxes. One word only.
        """
        try:
            llm_before = API_TRACKER.llm_snapshot()
            result = Settings.llm.complete(eval_prompt).text.strip().upper()
            API_TRACKER.record_llm_operation(
                "faithfulness_evaluation",
                llm_before,
            )
            return "SAFE" in result
        except Exception:
            return True 

    def ask(
        self,
        question: str,
        rag_mode: str = "auto",
        max_retries: int = 3,
    ) -> str:
        if self.vector_index is None:
            if not self.load_indexes():
                return (
                    "The knowledge base is not initialized. "
                    "Add documents and synchronize first."
                )

        if not self.lexical_index.built:
            self.rebuild_lexical_index()

        print(f"\n[QUESTION]\n{question}")
        query_start = time.perf_counter()

        profile = self.profile_query(
            question,
            rag_mode=rag_mode,
        )

        print(
            f"[RETRIEVAL STRATEGY] "
            f"{profile.mode} "
            f"({profile.reason})"
        )

        llm_before = API_TRACKER.llm_snapshot()

        if profile.mode == "summary":
            if self.engine is None:
                self.build_router()

            if self.engine is None:
                generated_answer = (
                    "The summary retrieval engine is unavailable."
                )
                retrieved_contexts = []
            else:
                response = self.engine.query(question)
                generated_answer = response.response
                retrieved_contexts = [
                    node.node.get_content()
                    for node in response.source_nodes
                ]
        else:
            retrieved_nodes, _ = self.retrieve_adaptively(
                question,
                profile,
            )
            retrieved_nodes = retrieved_nodes or []

            generated_answer = self.synthesize_answer(
                question,
                retrieved_nodes,
            )
            retrieved_contexts = [
                node_with_score.node.get_content()
                for node_with_score in retrieved_nodes
            ]

        API_TRACKER.record_llm_operation(
            "query_routing_and_answer",
            llm_before,
        )

        unified_context = "\n---\n".join(
            retrieved_contexts
        )

        attempt = 0
        while attempt < max_retries:
            attempt += 1

            is_faithful = self.evaluate_faithfulness(
                generated_answer,
                retrieved_contexts,
            )

            if is_faithful:
                query_seconds = (
                    time.perf_counter() - query_start
                )
                API_TRACKER.record_query(
                    1,
                    query_seconds,
                    len(retrieved_contexts),
                )
                API_TRACKER.print_query_usage(
                    llm_before
                )

                print(
                    f"\n[ANSWER] "
                    f"(Verified Factual on Attempt {attempt})"
                )
                print(generated_answer)
                print("=" * 70)
                return generated_answer

            print(
                f"[AUDIT WARNING] Attempt {attempt} "
                f"failed factuality check. "
                f"Running self-correction..."
            )

            correction_prompt = f"""
You previously generated an ANSWER that contained fabrications, inferences,
or assumptions not explicitly backed by the verified CONTEXT.
Rewrite the response completely.

CRITICAL RULES:
1. Rely ONLY on clear facts explicitly stated in CONTEXT.
2. Do NOT extrapolate, assume, or use outside knowledge.
3. If the context does not explicitly contain the answer, say:
   "The provided documentation does not contain this information."

VERIFIED CONTEXT:
\"\"\"{unified_context}\"\"\"

YOUR PREVIOUS ANSWER:
\"\"\"{generated_answer}\"\"\"

ORIGINAL USER QUESTION:
"{question}"

Provide the corrected, strictly factual answer.
"""
            try:
                correction_before = (
                    API_TRACKER.llm_snapshot()
                )
                generated_answer = (
                    Settings.llm.complete(
                        correction_prompt
                    ).text.strip()
                )
                API_TRACKER.record_llm_operation(
                    "self_correction",
                    correction_before,
                )
            except Exception as error:
                print(
                    f"--> [ERROR] Network drop during "
                    f"correction retry: {error}"
                )
                break

        query_seconds = (
            time.perf_counter() - query_start
        )
        API_TRACKER.record_query(
            1,
            query_seconds,
            len(retrieved_contexts),
        )
        API_TRACKER.print_query_usage(
            llm_before
        )

        print(
            "[GUARDRAIL BLOCK] Maximum self-correction "
            "retries reached. Output completely masked."
        )
        return (
            "I apologize, but I am unable to verify or prove "
            "that claim using the uploaded documentation."
        )

    def clear_storage(self):

        print(
            "--> [WIPE] "
            "Deleting RAG storage and container volumes..."
        )

        try:
            collection_name = "pipeline_collection"
            collections = self.db_client.get_collections().collections
            if any(c.name == collection_name for c in collections):
                self.db_client.delete_collection(collection_name=collection_name)
                print("--> [WIPE] Qdrant collection dropped from Docker bubble.")
        except Exception as e:
            print(f"--> [WARNING] Failed to drop Qdrant collection: {e}")


        if STORAGE_DIR.exists():

            for item in (
                STORAGE_DIR.iterdir()
            ):

                if item.is_file():

                    item.unlink()

                elif item.is_dir():

                    shutil.rmtree(
                        item
                    )

        STORAGE_DIR.mkdir(
            parents=True,
            exist_ok=True,
        )



if __name__ == "__main__":

    rag = AdaptiveRAG(
        data_dir="./data"
    )

    print("--> [SYSTEM LOG] Initiating file inventory and database sync...")


    rag.sync(
        force_rebuild=False
    )
    print("--> [SYSTEM LOG] Synchronized successfully. Knowledge base is online.")

    print("\n" + "=" * 60)
    print("CONTAINERIZED ENTERPRISE RAG KNOWLEDGE BASE CORE")
    print("  Type your questions below. Type 'exit' or 'quit' to close.")
    print("=" * 60 + "\n")

    while True:
        try:
            user_query = input("Enter your query: ").strip()
            
            if user_query.lower() in {"exit", "quit"}:
                print("\n--> [SYSTEM LOG] Shutting down connection layers. Goodbye.")
                break
                
            if not user_query:
                continue

            rag_mode = input(
                "RAG mode [auto/semantic/keyword/hybrid/summary] "
                "(default auto): "
            ).strip() or "auto"

            rag.ask(
                user_query,
                rag_mode=rag_mode,
            )
            API_TRACKER.print_summary()
            
        except KeyboardInterrupt:
            print("\n\n--> [SYSTEM LOG] System execution interrupted by user. Closing safely.")
            break
        except Exception as e:
            print(f"\n[RUNTIME ERROR]: An error occurred: {e}\n")
